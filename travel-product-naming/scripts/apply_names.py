# -*- coding: utf-8 -*-
"""把 names.json 写回 Excel（本次任务实战代码）

备份 .bak -> 新增「编号」「新命名」两列(保留原名称列) -> 写入 -> 去重校验

用法:
  python apply_names.py --excel 清单.xlsx --info products_info.json --names names.json --bak
可选:
  --no-col 11 --name-col 12
"""
import argparse, json, shutil
from collections import Counter
import openpyxl


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--excel', required=True, help='目标 xlsx')
    ap.add_argument('--info', default='products_info.json', help='ourtour_fetch.py 产物(含 row/no 映射)')
    ap.add_argument('--names', required=True, help='命名清单 json: {编号: 完整命名}')
    ap.add_argument('--no-col', type=int, default=11, help='编号写入列')
    ap.add_argument('--name-col', type=int, default=12, help='新命名写入列')
    ap.add_argument('--bak', action='store_true', help='写前备份为 .bak')
    args = ap.parse_args()

    info = json.load(open(args.info, encoding='utf-8'))
    rows = {it['no']: it['row'] for it in info}
    names = json.load(open(args.names, encoding='utf-8'))

    if args.bak:
        shutil.copy(args.excel, args.excel + '.bak')

    wb = openpyxl.load_workbook(args.excel)
    ws = wb[wb.sheetnames[0]]
    ws.cell(1, args.no_col, '编号')
    ws.cell(1, args.name_col, '新命名')
    n = 0
    for no, new in names.items():
        r = rows.get(no)
        if not r:
            print('NO ROW FOR', no)
            continue
        ws.cell(r, args.no_col, no)
        ws.cell(r, args.name_col, new)
        n += 1
    wb.save(args.excel)

    dup = [k for k, v in Counter(names.values()).items() if v > 1]
    print('written', n, '/', len(names))
    if dup:
        print('DUPLICATES(需回到原名称与参考文件找差异区分):', dup)
    else:
        print('去重校验通过')


if __name__ == '__main__':
    main()