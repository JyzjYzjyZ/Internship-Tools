# -*- coding: utf-8 -*-
"""联网抓取 ourtour 产品详情（本次任务实战代码）

读 Excel 清单 -> 按红黄绿灯分配编号(R/Y/G/W) -> 逐条调搜索接口抓详情 -> products_info.json

用法:
  python ourtour_fetch.py --excel 清单.xlsx --out products_info.json
可选:
  --sheet Sheet1 --start-row 2
  --col-name 2 --col-light 6 --col-pid 10
"""
import argparse, json, re, time, datetime
import urllib.request
import openpyxl

API = 'https://pro-api.ourtour.com/openapi/search/product'
LETTER = {'红灯': 'R', '黄灯': 'Y', '绿灯': 'G', None: 'W', '': 'W'}


def ser(o):
    if isinstance(o, (datetime.datetime, datetime.date)):
        return o.strftime('%Y-%m-%d')
    return o


def post_json(url, payload):
    req = urllib.request.Request(url, data=json.dumps(payload).encode('utf-8'),
                                 headers={'Content-Type': 'application/json;charset=utf-8',
                                          'User-Agent': 'Mozilla/5.0'})
    return json.loads(urllib.request.urlopen(req, timeout=30).read().decode('utf-8', 'ignore'))


def search(kw):
    try:
        resp = post_json(API, {'keyword': kw, 'pageNo': 1, 'pageSize': 5})
        recs = (resp.get('data') or {}).get('productQueryResultTOList', [{}])[0] \
                  .get('productQueryResultTOPage', {}).get('records', [])
        return recs
    except Exception as e:
        return [{'ERR': str(e)}]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--excel', required=True, help='产品清单 xlsx')
    ap.add_argument('--out', default='products_info.json', help='输出 json(UTF-8)')
    ap.add_argument('--sheet', default='Sheet1')
    ap.add_argument('--start-row', type=int, default=2, help='数据起始行(表头行不固定，先读前几行确认)')
    ap.add_argument('--col-name', type=int, default=2, help='产品名称列')
    ap.add_argument('--col-light', type=int, default=6, help='红黄绿灯列')
    ap.add_argument('--col-pid', type=int, default=10, help='产品编号列(可能混有"产品编号：xxx"字符串)')
    args = ap.parse_args()

    ws = openpyxl.load_workbook(args.excel)[args.sheet]
    items = []
    for r in range(args.start_row, ws.max_row + 1):
        name = ws.cell(r, args.col_name).value
        if not name:
            continue
        light = ws.cell(r, args.col_light).value
        pid = ws.cell(r, args.col_pid).value
        p = None
        if isinstance(pid, (int, float)):
            p = str(int(pid))
        elif isinstance(pid, str):
            m = re.search(r'(\d{6,})', pid)
            if m:
                p = m.group(1)
        items.append({'row': r, 'name': str(name).strip(), 'light': light, 'pid': p})

    seq = {k: 0 for k in LETTER}
    for it in items:
        seq[it['light']] += 1
        it['no'] = f"{LETTER[it['light']]}{seq[it['light']]:02d}"
    print('分组计数:', {k: v for k, v in seq.items()})

    for it in items:
        recs = None
        if it['pid']:
            recs = search(it['pid'])
            # 命中校验：返回的 productId 必须等于查询编号，或名称含表内关键字，
            # 否则可能是误匹配的无关产品(如门票编号搜到酒店套餐)，弃用
            ok = [x for x in recs
                  if str(x.get('productId')) == it['pid']
                  or it['name'][:4] in (x.get('productName') or '')]
            recs = ok or None
        if not recs:
            kw = re.sub(r'\s+', ' ', it['name'])[:30]
            recs = search(kw)
            if recs and 'ERR' in recs[0]:
                recs = []
        it['api'] = recs[0] if recs and recs[0].get('productId') else None
        time.sleep(0.3)  # 防限流

    with open(args.out, 'w', encoding='utf-8') as f:
        json.dump([{k: ser(v) for k, v in it.items()} for it in items], f,
                  ensure_ascii=False, indent=1)

    miss = [it for it in items if not it['api']]
    print('total', len(items), 'fetched', len(items) - len(miss))
    for it in miss:
        print('MISSING', it['no'], it['name'][:40], it['pid'], '-> 用表内信息命名')


if __name__ == '__main__':
    main()